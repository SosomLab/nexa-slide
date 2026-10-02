# -*- coding: utf-8 -*-
"""덱 JSON → 편집 가능한 PPTX (python-pptx, 모든 요소를 PowerPoint 기본 도형으로).

    python3 studio/export_pptx.py ch00            # → out/ch00.pptx
    python3 studio/export_pptx.py ch00 --out x.pptx

단위: 요소 좌표 px(1280×720) → EMU = px × 9525, 글자 size(px) → pt = px × 0.75.
줄 간격은 "정확히(pt)" = size × lineHeight × 0.75 로 넣어 브라우저 line-height 와 같게 한다.
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (cur_fonts, use_fonts, CODE_GUTTER_W, CODE_PAD_Y, OUT, REPO, HL_ROLE, code_lines, code_style, color, load_deck, plan_runs, radius,  # noqa: E402
                    runs, safe_id, table_grid, tokens)

from lxml import etree  # noqa: E402  (python-pptx 의존성)
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Emu  # noqa: E402

PX = 9525


def E(px):
    return Emu(int(round(float(px) * PX)))


def hpt(px):
    """px → 1/100 pt (글자 크기·줄 간격 XML 값)"""
    return int(round(float(px) * 0.75 * 100))


ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
NOSTYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid


def rgb_el(parent_tag, hexv):
    el = etree.SubElement(parent_tag, qn("a:srgbClr")) if parent_tag is not None else etree.Element(qn("a:srgbClr"))
    el.set("val", hexv.lstrip("#").upper())
    return el


def fonts_of(key):
    fs = cur_fonts()  # 슬라이드·덱 글꼴 세트(fontPreset)가 있으면 그 글꼴
    f = fs.get(key or "body", fs["body"])
    return f["latin"], f["ea"]


def set_run(r, text, size, bold=False, col="#000000", font="body", highlight=None, italic=False):
    """run 서식을 XML로 직접 쓴다(순서: solidFill → highlight → latin → ea → cs)."""
    r.text = text
    rPr = r._r.get_or_add_rPr()
    for c in list(rPr):
        rPr.remove(c)
    rPr.set("lang", "ko-KR")
    rPr.set("altLang", "en-US")
    rPr.set("sz", str(hpt(size)))
    rPr.set("b", "1" if bold else "0")
    rPr.set("i", "1" if italic else "0")
    rPr.set("dirty", "0")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    rgb_el(sf, col or "#000000")
    if highlight:
        hl = etree.SubElement(rPr, qn("a:highlight"))
        rgb_el(hl, highlight)
    latin, ea = fonts_of(font)
    etree.SubElement(rPr, qn("a:latin")).set("typeface", latin)
    etree.SubElement(rPr, qn("a:ea")).set("typeface", ea)
    etree.SubElement(rPr, qn("a:cs")).set("typeface", latin)


def end_para(p, size):
    """빈 줄도 같은 높이를 갖도록 endParaRPr 크기 지정"""
    epr = p._p.find(qn("a:endParaRPr"))
    if epr is None:
        epr = etree.SubElement(p._p, qn("a:endParaRPr"))
    epr.set("lang", "ko-KR")
    epr.set("sz", str(hpt(size)))


def para_fmt(p, size, lh, align="left", after_px=0):
    p.alignment = ALIGN.get(align, PP_ALIGN.LEFT)
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:lnSpc", "a:spcBef", "a:spcAft"):
        for c in pPr.findall(qn(tag)):
            pPr.remove(c)
    ln = etree.SubElement(pPr, qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPts")).set("val", str(hpt(size * lh)))
    sb = etree.SubElement(pPr, qn("a:spcBef"))
    etree.SubElement(sb, qn("a:spcPts")).set("val", "0")
    sa = etree.SubElement(pPr, qn("a:spcAft"))
    etree.SubElement(sa, qn("a:spcPts")).set("val", str(hpt(after_px)))
    return pPr


def frame(tf, pad=(0, 0, 0, 0), valign="top", wrap=True):
    tf.word_wrap = bool(wrap)
    tf.auto_size = MSO_AUTO_SIZE.NONE
    t, r, b, l = pad
    tf.margin_top, tf.margin_right, tf.margin_bottom, tf.margin_left = E(t), E(r), E(b), E(l)
    tf.vertical_anchor = ANCHOR.get(valign, MSO_ANCHOR.TOP)
    bp = tf._txBody.find(qn("a:bodyPr"))
    bp.set("rtlCol", "0")


def pad_of(el, default=0):
    p = el.get("pad", default)
    if isinstance(p, (int, float)):
        return (p, p, p, p)
    p = list(p) + [0] * 4
    return tuple(p[:4])


def fill_markup_text(tf, el, page, default_color="on-surface"):
    """text/pill/rect 글자(인라인 강조 포함) → 단락들"""
    size = float(el.get("size", 20))
    lh = float(el.get("lineHeight", 1.2 if el["type"] in ("pill", "ellipse", "rect") else 1.4))
    base = color(el.get("color", default_color))
    bold = bool(el.get("bold", el["type"] == "pill"))
    font = el.get("font", "body")
    lines = str(el.get("text", "")).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para_fmt(p, size, lh, el.get("align", "center" if el["type"] in ("pill", "ellipse") else "left"))
        for t, st in runs(line, page):
            if t:
                set_run(p.add_run(), t, size, st.get("bold", bold), color(st.get("color")) if st.get("color") else base, font)
        end_para(p, size)


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def shape_fill_line(shape, el):
    f = el.get("fill")
    if f:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color(f)[1:])
    else:
        shape.fill.background()
    if el.get("stroke"):
        shape.line.color.rgb = RGBColor.from_string(color(el["stroke"])[1:])
        shape.line.width = E(el.get("strokeWidth", 1))
    else:
        shape.line.fill.background()


def add_shape(slide, el, kind):
    x, y, w, h = el["x"], el["y"], el["w"], el["h"]
    if kind == "ellipse":
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x), E(y), E(w), E(h))
    else:
        r = min(w, h) / 2 if kind == "pill" else radius(el.get("radius"), w, h)
        if r > 0:
            shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
            shp.adjustments[0] = min(0.5, r / min(w, h))
        else:
            shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x), E(y), E(w), E(h))
    strip_style(shp)
    shape_fill_line(shp, el)
    return shp


def text_pad(el, kind, default=0):
    """PowerPoint 기본 도형의 글자 영역은 도형 안쪽으로 들어가 있다(둥근 사각형 = 반경 × 0.29289, 타원 = 폭·높이 × 0.14645).
    브라우저처럼 도형 가장자리에서 pad 만큼 떨어지도록 그만큼 여백에서 뺀다(음수는 0)."""
    t, r, b, l = pad_of(el, default)
    w, h = float(el["w"]), float(el["h"])
    if kind == "ellipse":
        ix, iy = w * 0.14645, h * 0.14645
    else:
        rad = min(w, h) / 2 if kind == "pill" else radius(el.get("radius"), w, h)
        ix = iy = 0.29289 * rad
    return (max(0, t - iy), max(0, r - ix), max(0, b - iy), max(0, l - ix))


# ---------------------------------------------------------------- 요소별
def do_box(slide, el, page):
    """rect · ellipse · pill · text(fill 있음/없음)"""
    t = el["type"]
    has_text = str(el.get("text", "")) != ""
    if t == "text" and not el.get("fill"):
        shp = slide.shapes.add_textbox(E(el["x"]), E(el["y"]), E(el["w"]), E(el["h"]))
    else:
        shp = add_shape(slide, el, "rect" if t == "text" else t)
    if has_text:
        tf = shp.text_frame
        default_valign = "top" if t == "text" else "middle"
        wrap = el.get("wrap", t in ("text", "rect"))
        pad = pad_of(el) if t == "text" and not el.get("fill") else text_pad(el, "rect" if t == "text" else t)
        frame(tf, pad, el.get("valign", default_valign), wrap)
        fill_markup_text(tf, el, page)
    elif t != "text":
        frame(shp.text_frame, (0, 0, 0, 0), "middle", False)
    return shp


def do_image(slide, el, page):
    p = REPO / el["src"]
    if not p.exists():
        print(f"  ! 이미지 없음: {el['src']}")
        return None
    return slide.shapes.add_picture(str(p), E(el["x"]), E(el["y"]), E(el["w"]), E(el["h"]))


def do_line(slide, el, page):
    x1, y1, x2, y2 = (el.get(k) for k in ("x1", "y1", "x2", "y2"))
    if x1 is None:
        x1, y1, x2, y2 = el["x"], el["y"], el["x"] + el["w"], el["y"] + el["h"]
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    strip_style(c)
    style_line(c.line._get_or_add_ln(), el)
    return c


def do_curve(slide, el, page):
    """곡선 — 꺾은선 근사 자유형(채우기 없음) + 끝 화살 머리. points 는 요소 상자 기준 px."""
    pts = [(el["x"] + x, el["y"] + y) for x, y in el.get("points", [])]
    if len(pts) < 2:
        return None
    fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
    shp = fb.convert_to_shape()
    strip_style(shp)
    shp.fill.background()
    style_line(shp.line._get_or_add_ln(), {**el, "type": "arrow" if el.get("head", "end") != "none" else "line"})
    return shp


def style_line(ln, el):
    for ch in list(ln):
        ln.remove(ch)
    ln.set("w", str(int(round(float(el.get("width", 2)) * PX))))
    ln.set("cap", "flat")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    rgb_el(sf, color(el.get("color", "on-surface")))
    if el.get("dash") and el["dash"] != "solid":
        etree.SubElement(ln, qn("a:prstDash")).set("val", "sysDot" if el["dash"] == "dot" else "dash")
    head = el.get("head", "end" if el["type"] == "arrow" else "none")
    if head in ("start", "both"):
        he = etree.SubElement(ln, qn("a:headEnd"))
        he.set("type", "triangle"), he.set("w", "lg"), he.set("len", "lg")
    if head in ("end", "both"):
        te = etree.SubElement(ln, qn("a:tailEnd"))
        te.set("type", "triangle"), te.set("w", "lg"), te.set("len", "lg")


def _cell_border(tcPr, side, col=None, w_px=1):
    ln = etree.SubElement(tcPr, qn(f"a:ln{side}"))
    if col:
        ln.set("w", str(int(w_px * PX)))
        ln.set("cap", "flat"), ln.set("cmpd", "sng"), ln.set("algn", "ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        rgb_el(sf, color(col))
        etree.SubElement(ln, qn("a:prstDash")).set("val", "solid")
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))


def do_table(slide, el, page):
    g = table_grid(el)
    nr, nc = len(g["cells"]), len(g["colW"])
    if not nr or not nc:
        return None
    gf = slide.shapes.add_table(nr, nc, E(g["x"]), E(g["y"]), E(g["tw"]), E(sum(g["rowH"])))
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        if a in tblPr.attrib:
            del tblPr.attrib[a]
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NOSTYLE_TABLE
    for j, cw in enumerate(g["colW"]):
        tbl.columns[j].width = E(cw)
    size = float(g["size"])
    for i in range(nr):
        tbl.rows[i].height = E(g["rowH"][i])
        for j in range(nc):
            c = g["cells"][i][j]
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = E(14)
            cell.margin_top = cell.margin_bottom = E(0)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            para_fmt(p, size, 1.4, c["align"])
            for t, st in runs(c["text"], page):
                if t:
                    set_run(p.add_run(), t, size, st.get("bold", c["bold"]),
                            color(st.get("color") or c["color"]), c["font"])
            end_para(p, size)
            tcPr = cell._tc.get_or_add_tcPr()
            for ch in list(tcPr):
                tcPr.remove(ch)
            tcPr.set("anchor", "ctr")
            _cell_border(tcPr, "L")
            _cell_border(tcPr, "R")
            _cell_border(tcPr, "T")
            is_head = i == 0 and (el.get("type") == "plan-table" or el.get("header", True))
            _cell_border(tcPr, "B", None if is_head else g["rule"], 1)
            if c["fill"]:
                sf = etree.SubElement(tcPr, qn("a:solidFill"))
                rgb_el(sf, color(c["fill"]))
            else:
                etree.SubElement(tcPr, qn("a:noFill"))
    shapes = [gf]
    for co in g["callouts"]:
        ln = {"type": "line", "x1": co["tx"] + 4, "y1": co["cy"], "x2": co["x"], "y2": co["cy"], "color": co["line"], "width": 2, "head": "none"}
        shapes.append(do_line(slide, ln, page))
        box = {"type": "text", "x": co["x"], "y": co["y"], "w": co["w"], "h": co["h"], "text": co["text"], "size": size,
               "lineHeight": 1.5, "fill": co["fill"], "color": co["color"], "radius": "r-s", "pad": [6, 14, 6, 14], "valign": "middle"}
        shapes.append(do_box(slide, box, page))
    return shapes


def do_plan(slide, el, page):
    """plan: 실행계획 원문(고정폭) + [[강조]] — run 글자색·굵게 + run 강조색(a:highlight)"""
    box = dict(el)
    box.setdefault("fill", "surface-container")
    box.setdefault("radius", "r-m")
    shp = add_shape(slide, box, "rect")
    tf = shp.text_frame
    frame(tf, text_pad(box, "rect", [16, 18, 16, 18]), "top", False)
    size = float(el.get("size", 14))
    lh = float(el.get("lineHeight", 1.5))
    base = color(el.get("color", "on-surface"))
    for i, line in enumerate(str(el.get("text", "")).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para_fmt(p, size, lh)
        for t, role in plan_runs(line):
            if not t:
                continue
            if role:
                bg, fg = HL_ROLE.get(role, HL_ROLE["warn"])
                set_run(p.add_run(), t, size, True, color(fg), "mono", highlight=color(bg))
            else:
                set_run(p.add_run(), t, size, False, base, "mono")
        end_para(p, size)
    return shp


def do_code(slide, el, page):
    """code: DBeaver 라이트 — 그룹 도형(여백 면·코드 면·강조 줄 면·테두리·줄 번호 글상자·코드 글상자).
    render.js elCode 와 같은 배치: 줄 번호 열 48px(오른쪽 여백 12), 코드 왼쪽 여백 14, 위아래 12."""
    x, y, w, h = float(el["x"]), float(el["y"]), float(el["w"]), float(el["h"])
    size = float(el.get("size", 15))
    lh = float(el.get("lineHeight", 1.6))
    L = size * lh
    gw = float(el.get("gutterW", CODE_GUTTER_W))
    r = radius(el.get("radius", "r-m"), w, h)
    lines = code_lines(el)
    grp = slide.shapes.add_group_shape()
    add_shape(grp, {"x": x, "y": y, "w": w, "h": h, "fill": "code-gutter-bg", "radius": r}, "rect")
    add_shape(grp, {"x": x + gw, "y": y, "w": w - gw, "h": h, "fill": "code-bg", "radius": r}, "rect")
    if r > 0:
        add_shape(grp, {"x": x + gw, "y": y, "w": min(r, w - gw), "h": h, "fill": "code-bg"}, "rect")
    for i, ln in enumerate(lines):
        if ln["focus"]:
            add_shape(grp, {"x": x + gw, "y": y + CODE_PAD_Y + i * L, "w": w - gw, "h": L, "fill": "code-focus"}, "rect")
    add_shape(grp, {"x": x + 0.5, "y": y + 0.5, "w": w - 1, "h": h - 1, "radius": max(0, r - 0.5),
                    "stroke": "surface-container-high", "strokeWidth": 1}, "rect")
    # 줄 번호
    g = grp.shapes.add_textbox(E(x), E(y + CODE_PAD_Y), E(gw - 12), E(len(lines) * L))
    frame(g.text_frame, (0, 0, 0, 0), "top", False)
    for i, ln in enumerate(lines):
        p = g.text_frame.paragraphs[0] if i == 0 else g.text_frame.add_paragraph()
        para_fmt(p, size, lh, "right")
        set_run(p.add_run(), ln["n"], size, ln["focus"], color("on-surface" if ln["focus"] else "code-gutter"), "mono")
        end_para(p, size)
    g.name = f"{el.get('id')} code-gutter"
    # 코드
    c = grp.shapes.add_textbox(E(x + gw + 14), E(y + CODE_PAD_Y), E(max(1, w - gw - 14)), E(len(lines) * L))
    frame(c.text_frame, (0, 0, 0, 0), "top", False)
    for i, ln in enumerate(lines):
        p = c.text_frame.paragraphs[0] if i == 0 else c.text_frame.add_paragraph()
        para_fmt(p, size, lh)
        for t, role in ln["runs"]:
            col, bold, italic = code_style(role)
            set_run(p.add_run(), t, size, bold, color(col), "mono", italic=italic)
        end_para(p, size)
    c.name = f"{el.get('id')} code-text"
    return grp


def bullet_rows(el):
    """bullets → [(글자, 수준, 뒤 간격 px)] — render.js bulletRows 와 같은 규칙"""
    gap, sub_gap = el.get("gap", 14), el.get("subGap", 6)
    rows = []
    for it in el.get("items", []):
        if isinstance(it, str):
            it = {"text": it}
        subs = it.get("sub", []) or []
        rows.append([it.get("text", ""), 0, sub_gap if subs else gap])
        for k, s in enumerate(subs):
            rows.append([s, 1, sub_gap if k < len(subs) - 1 else gap])
    return rows


def do_bullets(slide, el, page):
    shp = slide.shapes.add_textbox(E(el["x"]), E(el["y"]), E(el["w"]), E(el["h"]))
    tf = shp.text_frame
    frame(tf, pad_of(el), el.get("valign", "top"), True)
    size, sub = float(el.get("size", 20)), float(el.get("subSize", 17))
    lh = float(el.get("lineHeight", 1.55))
    for i, (t, lvl, after) in enumerate(bullet_rows(el)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sz = size if lvl == 0 else sub
        pPr = para_fmt(p, sz, lh, "left", after)
        marL, hang = (28, 28) if lvl == 0 else (48, 20)
        pPr.set("marL", str(marL * PX))
        pPr.set("indent", str(-hang * PX))
        pPr.set("lvl", "0")
        bc = etree.SubElement(pPr, qn("a:buClr"))
        rgb_el(bc, color(el.get("dot", "primary") if lvl == 0 else el.get("subColor", "on-surface-variant")))
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", "60000" if lvl == 0 else "100000")
        bf = etree.SubElement(pPr, qn("a:buFont"))
        bf.set("typeface", "맑은 고딕")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "●" if lvl == 0 else "–")
        base = color(el.get("color", "on-surface") if lvl == 0 else el.get("subColor", "on-surface-variant"))
        for tt, st in runs(t, page):
            if tt:
                set_run(p.add_run(), tt, sz, st.get("bold", False), color(st.get("color")) if st.get("color") else base, "body")
        end_para(p, sz)
    return shp


HANDLERS = {"rect": do_box, "ellipse": do_box, "pill": do_box, "text": do_box, "image": do_image,
            "line": do_line, "arrow": do_line, "curve": do_curve, "table": do_table, "plan-table": do_table,
            "code": do_code, "plan": do_plan, "bullets": do_bullets}


def export(deck, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(1280), E(720)
    blank = prs.slide_layouts[6]
    counts = []
    for n, s in enumerate(deck["slides"], 1):
        use_fonts(s.get("fontPreset") or deck.get("fontPreset"))  # 슬라이드 > 덱 > 작업 공간
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(color(s.get("bg", "surface"))[1:])
        els = sorted(enumerate(s.get("elements", [])), key=lambda t: (t[1].get("z", 0), t[0]))
        for _, el in els:
            h = HANDLERS.get(el.get("type"))
            if not h:
                print(f"  ! {s['id']}/{el.get('id')}: 모르는 요소 {el.get('type')}")
                continue
            res = h(slide, el, n)
            for shp in (res if isinstance(res, list) else [res]):
                if shp is not None:
                    shp.name = f"{el.get('id')} {el.get('type')}"
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
        counts.append(len(slide.shapes))
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return counts


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("id")
    ap.add_argument("--out")
    a = ap.parse_args()
    if not safe_id(a.id):
        sys.exit("잘못된 덱 id")
    out = Path(a.out) if a.out else OUT / f"{a.id}.pptx"
    counts = export(load_deck(a.id), out)
    print(f"{out} — 슬라이드 {len(counts)}장, 도형 {sum(counts)}개 ({', '.join(map(str, counts))})")
    if not a.out:  # 실제로 쓴 글꼴 기록 → out/<덱>.fonts.json (기록 실패는 내보내기를 막지 않는다)
        try:
            import font_report
            fo, rep = font_report.write(a.id, out)
            print(f"{fo} — 글꼴 " + ", ".join(u["typeface"] for u in rep["used"]))
        except Exception as e:  # noqa: BLE001
            print(f"글꼴 기록 실패: {e}", file=sys.stderr)


if __name__ == "__main__":
    main()
