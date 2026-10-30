# -*- coding: utf-8 -*-
"""PPT 파일로 시작(0.25.14) - 새로 만들기의 deckFrom → init_workspace --pptx-import → import_pptx --deck.

    python3 studio/tests/pptx_start_test.py

시험 안에서 작은 PPTX(제목·글머리·표)를 만들고, 사용자 설정은 임시 폴더로 격리한다(실제 최근 작업·첨부를 건드리지 않는다).
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
STUDIO = HERE.parent
WORK = Path(tempfile.mkdtemp(prefix="nexa-pptx-start-"))
os.environ["NEXA_SLIDE_USER_DIR"] = str(WORK / "user")
os.environ.pop("NEXA_SLIDE_WORKSPACE", None)
sys.path.insert(0, str(STUDIO))
FAILS = []


def check(name, cond, extra=""):
    print(("OK  " if cond else "FAIL"), name, "" if cond else extra)
    if not cond:
        FAILS.append(name)


def make_pptx(p):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "가져온 제목"
    s.placeholders[1].text_frame.text = "첫째 항목"
    s.placeholders[1].text_frame.add_paragraph().text = "둘째 항목"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "표 장"
    t = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "항목", "값", "가", "1"
    prs.save(str(p))


def main():
    import hub
    src = WORK / "원본 발표.pptx"
    make_pptx(src)
    hub.NEW_ATTACH.mkdir(parents=True, exist_ok=True)
    shutil.copy(src, hub.NEW_ATTACH / "ab12.pptx")  # /api/project attach 가 두는 이름 꼴

    # 1) 계획 - deckFrom 이면 --pptx-import, 시작용 내용은 빠진다
    base = {"title": "PPT 시작 시험", "folder": str(WORK / "proj"), "template": "lecture", "starter": "lecture-course", "kind": "deck"}
    p = hub.plan({**base, "deckFrom": {"name": "ab12.pptx", "orig": "원본 발표.pptx"}})
    a = p["args"]
    check("계획: --pptx-import <첨부> · --starter 없음", "--pptx-import" in a and a[a.index("--pptx-import") + 1].endswith("ab12.pptx") and "--starter" not in a, json.dumps(a, ensure_ascii=False))
    p2 = hub.plan({**base, "deckFrom": {"name": "../x.pptx"}})
    check("계획: 이상한 첨부 이름은 무시(시작용 내용 그대로)", "--pptx-import" not in p2["args"], json.dumps(p2["args"], ensure_ascii=False))
    p3 = hub.plan({**base, "kind": "sheet", "deckFrom": {"name": "ab12.pptx"}})
    check("계획: 시트에는 쓰지 않는다", "--pptx-import" not in p3["args"])

    # 2) 작업 공간 만들기 - 가져온 덱이 첫 덱, content 원본은 없다
    ws_root = WORK / "ws"
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    r = subprocess.run([sys.executable, str(STUDIO / "init_workspace.py"), str(ws_root), "--title", "PPT 시작", "--port", "5694", "--pptx-import", str(src), "--build", "--no-remember"],
                       env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    ws = ws_root / "slides"
    deck_p = ws / "decks" / "intro.json"
    check("만들기 성공", r.returncode == 0 and deck_p.is_file(), (r.stdout + r.stderr)[-600:])
    if deck_p.is_file():
        d = json.loads(deck_p.read_text(encoding="utf-8"))
        txt = json.dumps(d, ensure_ascii=False)
        check("덱: 2장 · 제목·글머리·표 글이 들어 있다 · 덱 제목 = 작업 공간 제목", len(d.get("slides", [])) == 2 and "가져온 제목" in txt and "둘째 항목" in txt and "항목" in txt and d.get("title") == "PPT 시작",
              txt[:300])
    check("content/intro.json(빈 시작 원본)은 만들지 않는다", not (ws / "content" / "intro.json").exists())
    check("출력에 빌드한 덱 intro", "빌드한 덱: intro" in r.stdout, r.stdout[-400:])
    r2 = subprocess.run([sys.executable, str(STUDIO / "init_workspace.py"), str(WORK / "ws2"), "--pptx-import", str(src), "--starter", "lecture-course", "--no-remember"],
                        env=env, capture_output=True, text=True, encoding="utf-8", errors="replace")
    check("시작용 교안과 함께 쓰면 거절", r2.returncode != 0 and "--pptx-import" in (r2.stderr + r2.stdout), (r2.stdout + r2.stderr)[-300:])

    shutil.rmtree(WORK, ignore_errors=True)
    print("\n실패:", FAILS or "없음")
    print("전부 통과" if not FAILS else "")
    sys.exit(1 if FAILS else 0)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    main()
