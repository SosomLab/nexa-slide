# -*- coding: utf-8 -*-
"""첨부 PPTX 내용 추출 - 세션이 Read 로 열 수 없는 PPTX(바이너리)에서 글·표·차트·발표자 노트를 JSON·Markdown 으로, 그림을 파일로 꺼낸다.

    python3 nexa.py import_pptx <파일.pptx> [--out <폴더>] [--no-images] [--json]
    → <폴더>/slides.md        슬라이드별 제목·본문(들여쓰기 = 수준)·표·차트 값·그림 경로·노트 - Read 로 읽어 덱을 쓴다
      <폴더>/slides.json      같은 내용의 구조(도형마다 kind·role·box[x,y,w,h] 1280×720 기준·pt = 가장 큰 글자 크기)
      <폴더>/images/sNN-MM.*  그림(같은 그림은 한 번만). 덱 image 요소의 src 에 이 경로를 그대로 쓴다
    기본 폴더: out/import/<파일 이름>/ (산출물 폴더, git 제외). --out 은 작업 공간 안이어야 한다.

    python3 nexa.py import_pptx --selftest [--out <폴더>]   # 견본 PPTX 를 만들어 추출을 확인(개발·설치 점검용)

꺼내는 것: 제목·부제·본문 자리(placeholder)와 일반 도형의 글(굵은 run 은 **…**), 표(행×열), 차트(종류·제목·범주·계열 값),
그림(원본 형식 그대로 - emf·wmf·bmp·tiff 는 Pillow 가 있으면 PNG 로), SmartArt 의 글, 그룹 안 도형(좌표는 슬라이드 기준으로 환산), 발표자 노트, 숨긴 슬라이드 표시.
PPTX 를 덱 JSON 으로 자동 변환하지는 않는다 - 세션이 내용을 읽고 slide_types·layout_samples 로 알맞은 장을 고른다.
"""
import argparse
import os
import base64
import hashlib
import io
import json
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import OUT, WORKSPACE  # noqa: E402

SLIDE_W = 1280  # 덱 좌표계(docs/deck-format.md) - PPTX 폭을 여기에 맞춰 환산
WEB_EXT = {"png", "jpg", "jpeg", "gif", "svg", "webp"}  # 편집기·PPTX 가 그대로 쓰는 형식 - 나머지는 PNG 로 바꿔 본다
# 1×1 투명 PNG(--selftest 의 그림용 - Pillow 없이도 견본을 만든다)
_PNG1 = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def _ns():
    from pptx.oxml.ns import qn
    return qn


# ---------- 좌표 - 그룹 안 도형은 그룹의 chOff/chExt → off/ext 로 환산해 슬라이드 기준 EMU 로 ----------
class _Xf:
    """EMU 좌표 변환(그룹 중첩 합성): x' = ox + (x - cx) * sx."""

    def __init__(self, ox=0, oy=0, cx=0, cy=0, sx=1.0, sy=1.0):
        self.ox, self.oy, self.cx, self.cy, self.sx, self.sy = ox, oy, cx, cy, sx, sy

    def pt(self, x, y):
        return self.ox + (x - self.cx) * self.sx, self.oy + (y - self.cy) * self.sy

    def box(self, sh):
        if sh.left is None or sh.width is None:
            return None
        x, y = self.pt(sh.left, sh.top)
        return [x, y, sh.width * self.sx, sh.height * self.sy]

    def child(self, grp):
        """그룹 안 좌표계 → 이 변환에 합성한 변환."""
        qn = _ns()
        xfrm = grp._element.find(qn("p:grpSpPr") + "/" + qn("a:xfrm"))
        if xfrm is None:
            return self
        off, ext, choff, chext = (xfrm.find(qn(t)) for t in ("a:off", "a:ext", "a:chOff", "a:chExt"))
        if off is None or ext is None:
            return self
        ox, oy = self.pt(int(off.get("x")), int(off.get("y")))
        cx, cy = (int(choff.get("x")), int(choff.get("y"))) if choff is not None else (int(off.get("x")), int(off.get("y")))
        cw = int(chext.get("cx")) if chext is not None and int(chext.get("cx")) else int(ext.get("cx"))
        ch = int(chext.get("cy")) if chext is not None and int(chext.get("cy")) else int(ext.get("cy"))
        sx = int(ext.get("cx")) / cw * self.sx if cw else self.sx
        sy = int(ext.get("cy")) / ch * self.sy if ch else self.sy
        return _Xf(ox, oy, cx, cy, sx, sy)


# ---------- 도형 → 항목 ----------
def _paragraphs(tf):
    """글 상자의 문단 → ([{level, text}], 가장 큰 글자 pt) - 빈 문단은 뺀다.
    굵은 run 은 **…**(덱의 인라인 강조와 같은 표기, 앞뒤 공백은 밖에), 줄바꿈(수직 탭)은 \\n. 글자 크기는 run 에 직접 적힌 것만 본다."""
    out, size = [], None
    for p in tf.paragraphs:
        parts = []
        for r in p.runs:
            t = r.text
            if not t:
                continue
            if r.font.size is not None:
                size = max(size or 0, r.font.size.pt)
            if r.font.bold and t.strip():
                core = t.strip()
                lead, trail = t[:len(t) - len(t.lstrip())], t[len(t.rstrip()):]
                t = f"{lead}**{core}**{trail}"
            parts.append(t)
        text = "".join(parts).replace("\x0b", "\n").strip() or p.text.replace("\x0b", "\n").strip()
        if text:
            out.append({"level": p.level or 0, "text": text})
    return out, size


def _role(sh):
    """자리(placeholder) 종류 → 역할 이름(title·subtitle·body·…). 자리가 아니면 None."""
    if not sh.is_placeholder:
        return None
    try:
        t = sh.placeholder_format.type
    except Exception:  # noqa: BLE001
        return "placeholder"
    name = str(getattr(t, "name", t) or "").lower()
    if name in ("title", "center_title", "vertical_title"):
        return "title"
    if name == "subtitle":
        return "subtitle"
    if name in ("body", "object", "vertical_body"):
        return "body"
    return name or "placeholder"


def _chart(sh):
    ch = sh.chart
    item = {"kind": "chart", "type": str(getattr(ch.chart_type, "name", ch.chart_type)), "title": None, "categories": [], "series": []}
    try:
        if ch.has_title:
            item["title"] = ch.chart_title.text_frame.text.strip() or None
    except Exception:  # noqa: BLE001
        pass
    for plot in ch.plots:
        try:
            cats = [str(c) if c is not None else "" for c in plot.categories]
        except Exception:  # noqa: BLE001 - 산점도·거품형은 범주가 없다
            cats = []
        if cats and not item["categories"]:
            item["categories"] = cats
        for s in plot.series:
            try:
                vals = [v for v in s.values]
            except Exception:  # noqa: BLE001
                vals = []
            item["series"].append({"name": s.name, "values": vals})
    return item


def _smartart_texts(sh):
    """SmartArt(도식) 데이터 파트의 글 - 도형은 가져오지 않고 글만."""
    qn = _ns()
    from lxml import etree
    rel = sh._element.find(".//" + qn("dgm:relIds"))
    if rel is None:
        return None
    rid = rel.get(qn("r:dm"))
    try:
        part = sh.part.related_part(rid)  # python-pptx 1.0
    except AttributeError:
        part = sh.part.related_parts[rid]  # noqa - 옛 python-pptx
    except KeyError:
        return []
    try:
        root = etree.fromstring(part.blob)
    except Exception:  # noqa: BLE001
        return []
    texts = []
    for pt in root.iter(qn("dgm:pt")):
        t = "".join(x.text or "" for x in pt.iter(qn("a:t"))).strip()
        if t:
            texts.append(t)
    return texts


def _is_diagram(sh):
    qn = _ns()
    gd = sh._element.find(".//" + qn("a:graphicData"))
    return gd is not None and "diagram" in (gd.get("uri") or "")


class _Images:
    """그림 저장 - 같은 내용(해시)은 한 파일. 웹 형식이 아니면 Pillow 로 PNG 변환을 시도한다."""

    def __init__(self, d, enabled=True):
        self.dir, self.enabled, self.seen, self.saved, self.skipped = d, enabled, {}, 0, []

    def add(self, img, sn, k):
        blob, ext = img.blob, (img.ext or "bin").lower()
        key = hashlib.sha1(blob).hexdigest()
        if key in self.seen:
            return self.seen[key]
        size = None
        if ext not in WEB_EXT or ext in ("jpg", "jpeg", "png"):
            try:
                from PIL import Image
                im = Image.open(io.BytesIO(blob))
                size = list(im.size)
                if ext not in WEB_EXT:
                    buf = io.BytesIO()
                    im.convert("RGBA").save(buf, format="PNG")
                    blob, ext = buf.getvalue(), "png"
            except Exception:  # noqa: BLE001 - Pillow 없음·변환 불가 → 원본 그대로
                pass
        name = f"s{sn:02d}-{k:02d}.{ext}"
        rec = {"file": f"images/{name}", "px": size, "ext": ext}
        if self.enabled:
            self.dir.mkdir(parents=True, exist_ok=True)
            (self.dir / name).write_bytes(blob)
            self.saved += 1
        if ext not in WEB_EXT:
            self.skipped.append(name)
        self.seen[key] = rec
        return rec


def _walk(shapes, xf, sn, images, scale, items):
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in shapes:
        st = sh.shape_type
        if st == MSO_SHAPE_TYPE.GROUP:
            _walk(sh.shapes, xf.child(sh), sn, images, scale, items)
            continue
        box = xf.box(sh)
        item = {"kind": "shape", "name": sh.name, "box": [round(v * scale) for v in box] if box else None}
        role = _role(sh)
        if role:
            item["role"] = role
        if getattr(sh, "has_table", False) and sh.has_table:
            item["kind"] = "table"
            item["rows"] = [[c.text.replace("\x0b", "\n").strip() for c in row.cells] for row in sh.table.rows]
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            item.update(_chart(sh))
        elif st == MSO_SHAPE_TYPE.PICTURE or (sh.is_placeholder and hasattr(sh, "image") and role == "picture"):
            try:
                rec = images.add(sh.image, sn, sum(1 for i in items if i["kind"] == "image") + 1)
            except Exception as e:  # noqa: BLE001 - 연결된(외부) 그림 등
                item.update({"kind": "image", "file": None, "note": f"그림을 꺼내지 못함: {e}"})
            else:
                item.update({"kind": "image", "file": rec["file"], "px": rec["px"]})
                crop = [getattr(sh, f"crop_{k}", 0) or 0 for k in ("left", "top", "right", "bottom")]
                if any(abs(c) > 0.001 for c in crop):
                    item["crop"] = [round(c, 3) for c in crop]
        elif _is_diagram(sh):
            item["kind"] = "smartart"
            item["texts"] = _smartart_texts(sh) or []
        elif st == MSO_SHAPE_TYPE.MEDIA:
            item["kind"] = "media"
        elif st in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
            item["kind"] = "ole"
        if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
            paras, size = _paragraphs(sh.text_frame)
            if paras:
                if item["kind"] == "shape":
                    item["kind"] = "text"
                item["paragraphs"] = paras
                if size:
                    item["pt"] = size
        if item["kind"] == "shape" and not item.get("role"):
            continue  # 글도 그림도 없는 장식 도형(선·사각형)은 뺀다
        items.append(item)


def extract(src, out_dir, images=True):
    """PPTX → {source, slideSize, count, slides:[…]} 를 돌려주고 그림을 out_dir/images/ 에 쓴다."""
    from pptx import Presentation
    prs = Presentation(str(src))
    scale = SLIDE_W / prs.slide_width if prs.slide_width else 1 / 9525
    imgs = _Images(out_dir / "images", images)
    doc = {"source": str(src), "slideSize": [SLIDE_W, round(prs.slide_height * scale)], "count": len(prs.slides), "slides": []}
    for n, sld in enumerate(prs.slides, 1):
        items = []
        _walk(sld.shapes, _Xf(), n, imgs, scale, items)
        title = None
        try:
            if sld.shapes.title is not None:
                title = sld.shapes.title.text_frame.text.replace("\x0b", " ").strip() or None
        except Exception:  # noqa: BLE001
            title = None
        if not title:
            t = next((i for i in items if i.get("role") == "title" and i.get("paragraphs")), None)
            if t is None:  # 제목 자리가 없는 덱(글 상자만) - 위쪽 1/3 안에서 글자가 가장 큰 글 상자를 제목으로 추정
                top = [i for i in items if i["kind"] == "text" and i.get("pt") and i.get("box") and i["box"][1] < doc["slideSize"][1] / 3]
                t = max(top, key=lambda i: i["pt"], default=None)
                if t is not None:
                    t["role"] = "title?"
            title = " ".join(p["text"] for p in t["paragraphs"]).replace("**", "") if t else None
        notes = None
        try:
            if sld.has_notes_slide and sld.notes_slide.notes_text_frame is not None:
                notes = sld.notes_slide.notes_text_frame.text.replace("\x0b", "\n").strip() or None
        except Exception:  # noqa: BLE001
            notes = None
        s = {"n": n, "layout": sld.slide_layout.name if sld.slide_layout is not None else None, "title": title, "shapes": items, "notes": notes}
        if sld._element.get("show") == "0":
            s["hidden"] = True
        doc["slides"].append(s)
    doc["images"] = {"saved": imgs.saved, "unconverted": imgs.skipped}
    return doc


def _md_cell(t):
    return t.replace("|", "\\|").replace("\n", " ")


def to_markdown(doc):
    L = [f"# {Path(doc['source']).name} - 슬라이드 {doc['count']}장", "",
         "도형 box 는 [x, y, w, h] (1280×720 기준). 그림은 images/ 아래 - 덱 image 요소의 src 에 그 경로를 쓴다.", ""]
    for s in doc["slides"]:
        head = f"## 슬라이드 {s['n']}" + (f" - {s['title']}" if s.get("title") else "") + (" (숨김)" if s.get("hidden") else "")
        L += [head, f"- 레이아웃: {s.get('layout') or '-'}"]
        for it in s["shapes"]:
            box = f" box {it['box']}" if it.get("box") else ""
            k = it["kind"]
            if k == "text":
                role = it.get("role") or it.get("name") or "글"
                L.append(f"- **{role}**" + (f" {it['pt']:g}pt" if it.get("pt") else "") + box)
                for p in it["paragraphs"]:
                    L.append("  " * (p["level"] + 1) + "- " + p["text"].replace("\n", "  \n" + "  " * (p["level"] + 2)))
            elif k == "table":
                rows = it["rows"]
                L.append(f"- **표** {len(rows)}×{len(rows[0]) if rows else 0}{box}")
                if rows:
                    L.append("")
                    L.append("  | " + " | ".join(_md_cell(c) for c in rows[0]) + " |")
                    L.append("  |" + "---|" * len(rows[0]))
                    for r in rows[1:]:
                        L.append("  | " + " | ".join(_md_cell(c) for c in r) + " |")
                    L.append("")
            elif k == "chart":
                L.append(f"- **차트** {it['type']}" + (f" - {it['title']}" if it.get("title") else "") + box)
                if it.get("categories"):
                    L.append("  - 범주: " + ", ".join(it["categories"]))
                for se in it["series"]:
                    L.append(f"  - 계열 {se['name'] or '(이름 없음)'}: " + ", ".join("" if v is None else str(v) for v in se["values"]))
                if it.get("paragraphs"):
                    L.append("  - 글: " + " / ".join(p["text"] for p in it["paragraphs"]))
            elif k == "image":
                px = f" {it['px'][0]}×{it['px'][1]}px" if it.get("px") else ""
                L.append(f"- **그림** {it.get('file') or '(없음)'}{px}{box}" + (f" - {it['note']}" if it.get("note") else "")
                         + (f" - 잘라냄 {it['crop']}" if it.get("crop") else ""))
            elif k == "smartart":
                L.append(f"- **SmartArt(도식)**{box}: " + (" / ".join(it["texts"]) if it["texts"] else "(글 없음)"))
            elif k in ("media", "ole"):
                L.append(f"- **{'미디어' if k == 'media' else '포함 개체(OLE)'}** {it.get('name')}{box} - 꺼내지 않음")
                if it.get("paragraphs"):
                    L.append("  - 글: " + " / ".join(p["text"] for p in it["paragraphs"]))
            else:
                L.append(f"- {it.get('role') or k} {it.get('name')}{box}")
        if s.get("notes"):
            L.append("- **발표자 노트**: " + s["notes"].replace("\n", "  \n  "))
        L.append("")
    im = doc.get("images") or {}
    if im.get("unconverted"):
        L.append(f"그림 중 PNG 로 바꾸지 못한 것(원본 형식 그대로, 편집기에 보이지 않을 수 있음): {', '.join(im['unconverted'])}")
    return "\n".join(L).rstrip() + "\n"


def _rel(p):
    """cwd(작업 공간) 기준 상대 경로 - 명령·덱 src 에 그대로 쓰기 쉽게."""
    try:
        return p.resolve().relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return p.as_posix()


def _allowed_roots():
    """읽을 수 있는 폴더 - 작업 공간, 웹 세션이 저장소 루트 접근을 받았으면 그 루트(NEXA_SLIDE_REPO_ROOT). 이 명령은 nexa.py 로 세션이 부르므로
    세션의 읽기 범위(작업 디렉터리 + --add-dir) 밖 PPTX 를 이 명령으로 열지 못하게 같은 범위로 막는다."""
    roots = [WORKSPACE.resolve()]
    if os.environ.get("NEXA_SLIDE_REPO_ROOT"):
        roots.append(Path(os.environ["NEXA_SLIDE_REPO_ROOT"]).resolve())
    return roots


def _inside(p, roots):
    p = Path(p).resolve()
    return any(p == r or r in p.parents for r in roots)


def run(src, out_dir, images=True, as_json=False):
    src = Path(src)
    if not _inside(src, _allowed_roots()):
        sys.exit(f"작업 공간 안의 파일만 열 수 있다: {src} - 요청 첨부(out/attach/<덱>/…)나 작업 공간 폴더에 둔 PPTX 를 준다")
    if not src.is_file():
        sys.exit(f"파일 없음: {src}")
    if not zipfile.is_zipfile(src):
        sys.exit(f"PPTX(zip) 가 아니다: {src} - .ppt(옛 형식)·PDF 는 PowerPoint 에서 PPTX 로 저장해 다시 올린다")
    out_dir = Path(out_dir)
    ws = WORKSPACE.resolve()
    if ws != out_dir.resolve() and ws not in out_dir.resolve().parents:
        sys.exit(f"--out 은 작업 공간({ws}) 안이어야 한다: {out_dir}")
    try:
        doc = extract(src, out_dir, images)
    except Exception as e:  # noqa: BLE001 - 손상된 파일·암호 등
        sys.exit(f"PPTX 를 열지 못했다: {e}")
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "slides.json").write_text(json.dumps(doc, ensure_ascii=False, indent=1) + "\n", encoding="utf-8", newline="\n")
    (out_dir / "slides.md").write_text(to_markdown(doc), encoding="utf-8", newline="\n")
    kinds = {}
    for s in doc["slides"]:
        for it in s["shapes"]:
            kinds[it["kind"]] = kinds.get(it["kind"], 0) + 1
    notes = sum(1 for s in doc["slides"] if s.get("notes"))
    if as_json:
        print(json.dumps({"ok": True, "out": _rel(out_dir), "md": _rel(out_dir / "slides.md"), "json": _rel(out_dir / "slides.json"),
                          "slides": doc["count"], "kinds": kinds, "notes": notes, "images": doc["images"]}, ensure_ascii=False))
        return doc
    print(f"추출: {_rel(out_dir / 'slides.md')}  (슬라이드 {doc['count']}장 · "
          + " · ".join(f"{k} {v}" for k, v in sorted(kinds.items())) + f" · 노트 {notes}장 · 그림 파일 {doc['images']['saved']}개)")
    print(f"구조: {_rel(out_dir / 'slides.json')}" + (f" · 그림: {_rel(out_dir / 'images')}/" if doc["images"]["saved"] else ""))
    if doc["images"]["unconverted"]:
        print("알림: PNG 로 바꾸지 못한 그림(원본 형식 그대로): " + ", ".join(doc["images"]["unconverted"]))
    print("다음: slides.md 를 Read 로 읽고 slide_types·layout_samples 로 장을 골라 덱을 쓴다. 그림은 image 요소 src 에 위 경로 그대로.")
    return doc


# ---------- 견본 PPTX 로 자기 점검 ----------
def _make_sample(path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Emu, Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "견본 강의안"
    s1.placeholders[1].text = "부제 - import_pptx 점검"
    s1.notes_slide.notes_text_frame.text = "첫 장 노트"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "본문과 표"
    tf = s2.placeholders[1].text_frame
    tf.text = "첫째 항목"
    r = tf.paragraphs[0].add_run()
    r.text = " 굵게"
    r.font.bold = True
    p2 = tf.add_paragraph()
    p2.text, p2.level = "둘째 수준", 1
    tbl = s2.shapes.add_table(2, 3, Inches(1), Inches(4), Inches(6), Inches(1.5)).table
    for i, t in enumerate(("항목", "값", "비고")):
        tbl.cell(0, i).text = t
    for i, t in enumerate(("A", "1", "x|y")):
        tbl.cell(1, i).text = t
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "차트·그림·그룹"
    cd = CategoryChartData()
    cd.categories = ["1분기", "2분기"]
    cd.add_series("매출", (10.5, 20))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(5), Inches(3), cd)
    s3.shapes.add_picture(io.BytesIO(_PNG1), Inches(6), Inches(1.5), Inches(2), Inches(2))
    grp = s3.shapes.add_group_shape()
    tb = grp.shapes.add_textbox(Inches(6), Inches(4), Inches(3), Inches(0.5))
    tb.text_frame.text = "그룹 안 글"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    prs.save(str(path))


def selftest(out_dir):
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    sample = out_dir / "sample.pptx"
    _make_sample(sample)
    doc = run(sample, out_dir, as_json=False)
    s1, s2, s3 = doc["slides"]
    kinds3 = [i["kind"] for i in s3["shapes"]]
    checks = [
        ("제목", s1["title"] == "견본 강의안"),
        ("노트", s1["notes"] == "첫 장 노트"),
        ("굵게·수준", any(p["text"] == "첫째 항목 **굵게**" for i in s2["shapes"] for p in i.get("paragraphs", []))
         and any(p["level"] == 1 and p["text"] == "둘째 수준" for i in s2["shapes"] for p in i.get("paragraphs", []))),
        ("표", any(i["kind"] == "table" and i["rows"] == [["항목", "값", "비고"], ["A", "1", "x|y"]] for i in s2["shapes"])),
        ("차트", any(i["kind"] == "chart" and i["categories"] == ["1분기", "2분기"] and i["series"][0]["values"] == [10.5, 20.0] for i in s3["shapes"])),
        ("그림", any(i["kind"] == "image" and (out_dir / i["file"]).is_file() for i in s3["shapes"])),
        ("그룹 글", any(i["kind"] == "text" and i["paragraphs"][0]["text"] == "그룹 안 글" and i["box"] == [576, 384, 288, 48] and i.get("pt") == 14
                     for i in s3["shapes"])),
        ("종류", "chart" in kinds3 and "image" in kinds3 and "text" in kinds3),
    ]
    bad = [n for n, ok in checks if not ok]
    print("selftest: " + ("모두 통과 " + str([n for n, _ in checks]) if not bad else "실패 " + str(bad)))
    return 0 if not bad else 1


def main():
    for st in (sys.stdout, sys.stderr):
        try:
            st.reconfigure(encoding="utf-8", errors="replace")
        except Exception:  # noqa: BLE001
            pass
    ap = argparse.ArgumentParser(description="PPTX 내용 추출 → slides.md · slides.json · images/ (작업 공간 out/import/<이름>/)")
    ap.add_argument("file", nargs="?", help="PPTX 파일(요청 첨부 out/attach/<덱>/….pptx 등)")
    ap.add_argument("--out", help="결과 폴더(작업 공간 안). 기본 out/import/<파일 이름>/")
    ap.add_argument("--no-images", action="store_true", help="그림 파일을 쓰지 않는다(목록만)")
    ap.add_argument("--json", action="store_true", help="요약을 JSON 한 줄로")
    ap.add_argument("--selftest", action="store_true", help="견본 PPTX 를 만들어 추출을 점검")
    a = ap.parse_args()
    if a.selftest:
        sys.exit(selftest(a.out or (OUT / ".studio" / "import-selftest")))
    if not a.file:
        ap.error("PPTX 파일을 준다")
    src = Path(a.file)
    out = Path(a.out) if a.out else OUT / "import" / src.stem
    run(src, out, images=not a.no_images, as_json=a.json)


if __name__ == "__main__":
    main()
